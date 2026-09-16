"""Standalone parser entry point. Never import application configuration/keys.

Run via isolated Python with bounded stdin/stdout. This limits resource use;
the worker's OS identity/filesystem/network still need deployment isolation.
"""
from io import BytesIO
import json
from pathlib import Path
import sys

MAX_INPUT_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 2 * 1024 * 1024
MAX_PAGES = 256
MEMORY_BYTES = 512 * 1024 * 1024
CPU_SECONDS = 10


def apply_limits():
    import resource
    limits = [(resource.RLIMIT_CPU, CPU_SECONDS), (resource.RLIMIT_CORE, 0),
              (resource.RLIMIT_FSIZE, 0)]
    # Linux enforces the entire virtual address space. macOS does not enforce
    # RLIMIT_AS and rejects useful RLIMIT_DATA values: CPU/time/I/O bounds only
    # there. Production memory isolation requires Linux (or an external sandbox).
    if sys.platform.startswith('linux'):
        limits.append((resource.RLIMIT_AS, MEMORY_BYTES))
    for kind, value in limits:
        _, hard = resource.getrlimit(kind)
        value = value if hard == resource.RLIM_INFINITY else min(value, hard)
        resource.setrlimit(kind, (value, value))


def parse_document(kind, content):
    if len(content) > MAX_INPUT_BYTES:
        raise ValueError('document input budget exceeded')
    parts = []
    total = 0
    def append(text):
        nonlocal total
        total += len(text) + 2
        if total > MAX_TEXT_CHARS:
            raise ValueError('document text budget exceeded')
        parts.append(text)

    if kind == 'docx':
        # Load the pure archive guard by absolute path, avoiding core.__init__
        # (which loads application configuration and logging).
        import runpy
        guard = runpy.run_path(str(Path(__file__).resolve().parents[1] / 'core/security/document_archive.py'))
        guard['validate_docx_archive'](content)
        from docx import Document
        doc = Document(BytesIO(content))
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        append(cell.text.strip())
        return {'content': '\n'.join(parts)}
    if kind != 'pdf':
        raise ValueError('unsupported document type')
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(content), strict=False)
    count = len(reader.pages)
    if count > MAX_PAGES:
        raise ValueError('document page budget exceeded')
    for page in reader.pages:
        append(page.extract_text() or '')
    metadata = {}
    for key, value in (reader.metadata or {}).items():
        if len(metadata) >= 64:
            break
        metadata[str(key).lstrip('/')[:128]] = str(value)[:4096]
    metadata['total_pages'] = count
    return {
        'content': '\n\n'.join(parts),
        'title': metadata.get('Title') or 'Untitled PDF Document',
        'pages': [{'number': idx + 1, 'content': text} for idx, text in enumerate(parts)],
        'metadata': metadata,
        'stats': {'original_size': len(content), 'processed_size': total,
                  'successful_pages': sum(bool(text.strip()) for text in parts),
                  'total_pages': count},
    }


def main():
    try:
        apply_limits()  # Fail before third-party imports/input if limits fail.
        content = sys.stdin.buffer.read(MAX_INPUT_BYTES + 1)
        result = parse_document(sys.argv[1], content)
        sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode('utf-8'))
    except BaseException:
        # Parser exceptions may contain document contents: do not log them.
        sys.stderr.write('document parsing refused or resource limit reached\n')
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
